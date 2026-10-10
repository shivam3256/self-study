import io
import pytest
from httpx import AsyncClient, ASGITransport
import openpyxl
from sqlalchemy import select

from app.main import app
from app.core.database import AsyncSessionLocal
from app.models.student import Student


def create_excel_file(rows: list, sheet_name="Students") -> bytes:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = sheet_name

    headers = [
        "Name",
        "Mob No.",
        "Gender",
        "ID Proof Type",
        "ID Proof",
        "Email",
        "Notes",
    ]
    ws.append(headers)

    for row in rows:
        ws.append(row)

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return output.getvalue()


@pytest.mark.asyncio
async def test_download_excel_template():
    transport = ASGITransport(app=app)
    async with AsyncClient(transport=transport, base_url="http://test") as ac:
        login_resp = await ac.post("/api/v1/auth/login", json={
            "email": "owner@apexlibrary.com",
            "password": "admin123"
        })
        token = login_resp.json()["access_token"]
        headers = {"Authorization": f"Bearer {token}"}

        resp = await ac.get("/api/v1/students/import/template", headers=headers)
        assert resp.status_code == 200
        assert resp.headers["content-type"] == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        
        # Verify content with openpyxl
        wb = openpyxl.load_workbook(io.BytesIO(resp.content))
        assert "Students" in wb.sheetnames
        assert "Instructions" in wb.sheetnames

        ws = wb["Students"]
        headers_in_sheet = [cell.value for cell in ws[1]]
        assert headers_in_sheet == [
            "Name", "Mob No.", "Gender", "ID Proof Type", "ID Proof", "Email", "Notes"
        ]

        # Verify Row 2 is the reference example (Rahul Sharma)
        assert ws.cell(row=2, column=1).value == "Rahul Sharma"
        assert ws.cell(row=2, column=3).value == "Male"


@pytest.mark.asyncio
async def test_ignore_template_example_row_on_import():
    transport = ASGITransport(app=app)
    async with AsyncClient(transport=transport, base_url="http://test") as ac:
        login_resp = await ac.post("/api/v1/auth/login", json={
            "email": "owner@apexlibrary.com",
            "password": "admin123"
        })
        token = login_resp.json()["access_token"]
        tenant_id = login_resp.json()["user"]["tenant_id"]
        headers = {"Authorization": f"Bearer {token}"}

        async with AsyncSessionLocal() as db:
            prev_st = (await db.execute(
                select(Student).where(Student.tenant_id == tenant_id, Student.phone == "+91 9123456789")
            )).scalars().all()
            for s in prev_st:
                await db.delete(s)
            await db.commit()

        # Create file with Row 1 = Example "Rahul Sharma", Row 2 = Real Student
        rows = [
            ["Rahul Sharma", "+91 9876543210", "Male", "Aadhaar", "1234-5678-9012", "rahul.sharma@example.com", "Civil services"],
            ["Real Actual Student", "+91 9123456789", "Female", "Aadhaar", "9876-5432-1098", "real.student@gmail.com", "UPSC Aspirant"]
        ]
        file_bytes = create_excel_file(rows)

        val_resp = await ac.post(
            "/api/v1/students/import/validate",
            files={"file": ("test_import_with_example.xlsx", file_bytes, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
            headers=headers
        )
        assert val_resp.status_code == 200
        val_data = val_resp.json()
        assert val_data["total_rows"] == 1
        assert val_data["valid_rows"] == 1
        assert val_data["error_rows"] == 0
        assert val_data["rows"][0]["name"] == "Real Actual Student"
        assert val_data["rows"][0]["gender"] == "female"

        # Confirm import
        confirm_resp = await ac.post(
            "/api/v1/students/import/confirm",
            json={"rows": val_data["rows"]},
            headers=headers
        )
        assert confirm_resp.status_code == 200
        assert confirm_resp.json()["imported_count"] == 1


@pytest.mark.asyncio
async def test_validate_and_confirm_valid_import():
    transport = ASGITransport(app=app)
    async with AsyncClient(transport=transport, base_url="http://test") as ac:
        login_resp = await ac.post("/api/v1/auth/login", json={
            "email": "owner@apexlibrary.com",
            "password": "admin123"
        })
        token = login_resp.json()["access_token"]
        tenant_id = login_resp.json()["user"]["tenant_id"]
        headers = {"Authorization": f"Bearer {token}"}

        async with AsyncSessionLocal() as db:
            prev_sts = (await db.execute(
                select(Student).where(Student.tenant_id == tenant_id, Student.phone.in_(["+91 9911992288", "9876599999"]))
            )).scalars().all()
            for s in prev_sts:
                await db.delete(s)
            await db.commit()

        # Create valid rows matching registration details
        rows = [
            ["Imported Student 1", "+91 9911992288", "Male", "Aadhaar", "9988-7766-5544", "student1@example.com", "CA Final"],
            ["Only Name Student", None, None, None, None, None, None],
            ["Imported Student 2", "9876599999", "Female", "College ID", "COL-9988", "student2@example.com", "NEET Prep"]
        ]
        file_bytes = create_excel_file(rows)

        # Step 1: Validate
        val_resp = await ac.post(
            "/api/v1/students/import/validate",
            files={"file": ("test_students.xlsx", file_bytes, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
            headers=headers
        )
        assert val_resp.status_code == 200
        val_data = val_resp.json()
        assert val_data["total_rows"] == 3
        assert val_data["valid_rows"] == 3
        assert val_data["error_rows"] == 0
        assert val_data["can_import"] is True

        # Step 2: Confirm
        confirm_resp = await ac.post(
            "/api/v1/students/import/confirm",
            json={"rows": val_data["rows"]},
            headers=headers
        )
        assert confirm_resp.status_code == 200
        res_data = confirm_resp.json()
        assert res_data["imported_count"] == 3
        assert res_data["skipped_count"] == 0


@pytest.mark.asyncio
async def test_validation_errors_missing_name_invalid_email():
    transport = ASGITransport(app=app)
    async with AsyncClient(transport=transport, base_url="http://test") as ac:
        login_resp = await ac.post("/api/v1/auth/login", json={
            "email": "owner@apexlibrary.com",
            "password": "admin123"
        })
        token = login_resp.json()["access_token"]
        headers = {"Authorization": f"Bearer {token}"}

        # Row 1: Missing Name
        # Row 2: Invalid Email
        rows = [
            ["   ", "+91 9911223344", "Male", "Aadhaar", None, None, None],
            ["Test Bad Student", "9876543210", "Female", "Aadhaar", None, "not-an-email", "Some note"]
        ]
        file_bytes = create_excel_file(rows)

        val_resp = await ac.post(
            "/api/v1/students/import/validate",
            files={"file": ("invalid_students.xlsx", file_bytes, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
            headers=headers
        )
        assert val_resp.status_code == 200
        val_data = val_resp.json()
        assert val_data["total_rows"] == 2
        assert val_data["error_rows"] == 2
        assert val_data["can_import"] is False

        row1 = val_data["rows"][0]
        assert not row1["is_valid"]
        assert any("Name is required" in err for err in row1["errors"])

        row2 = val_data["rows"][1]
        assert not row2["is_valid"]
        assert any("email" in err.lower() for err in row2["errors"])


@pytest.mark.asyncio
async def test_gender_normalization_in_import():
    transport = ASGITransport(app=app)
    async with AsyncClient(transport=transport, base_url="http://test") as ac:
        login_resp = await ac.post("/api/v1/auth/login", json={
            "email": "owner@apexlibrary.com",
            "password": "admin123"
        })
        token = login_resp.json()["access_token"]
        headers = {"Authorization": f"Bearer {token}"}

        rows = [
            ["Student M", "9000000001", "M", "Aadhaar", None, None, None],
            ["Student F", "9000000002", "Female", "Aadhaar", None, None, None],
            ["Student Other", "9000000003", "Non-binary", "Aadhaar", None, None, None],
        ]
        file_bytes = create_excel_file(rows)

        val_resp = await ac.post(
            "/api/v1/students/import/validate",
            files={"file": ("gender_test.xlsx", file_bytes, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
            headers=headers
        )
        assert val_resp.status_code == 200
        val_data = val_resp.json()
        assert val_data["valid_rows"] == 3
        assert val_data["rows"][0]["gender"] == "male"
        assert val_data["rows"][1]["gender"] == "female"
        assert val_data["rows"][2]["gender"] == "other"
